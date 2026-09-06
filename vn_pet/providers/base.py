from __future__ import annotations

import json
import base64
from abc import ABC, abstractmethod
from collections.abc import AsyncIterator
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Callable

from fastapi import HTTPException

from vn_pet.secrets import has_real_model_api_key
from vn_pet.secrets import EmptyModelResponseError


SystemPromptProvider = Callable[[], str]


@dataclass(frozen=True)
class LLMCallResult:
    text: str
    prompt_tokens: int = 0
    completion_tokens: int = 0
    total_tokens: int = 0
    model: str = ""
    usage_estimated: bool = True
    tool_calls: list[dict] | None = None
    reasoning_content: str | None = None
    finish_reason: str | None = None


@dataclass(frozen=True)
class LLMStreamChunk:
    delta: str = ""
    result: LLMCallResult | None = None
    tool_calls: list[dict] | None = None
    reasoning_delta: str = ""


@dataclass(frozen=True)
class ModelImage:
    name: str
    mime_type: str
    data: bytes


@dataclass(frozen=True)
class ModelImagePayload:
    """Host-internal image payload; never serialize this object into history."""

    public_text: str
    images: tuple[ModelImage, ...]


class ContextWindowExceeded(HTTPException):
    def __init__(self, detail: str = "模型上下文长度已超限") -> None:
        super().__init__(status_code=413, detail=detail)


class EmptyModelResponse(HTTPException, EmptyModelResponseError):
    """Raised when a model completes without text or tool calls."""

    def __init__(self, detail: str = "模型没有返回有效文本") -> None:
        super().__init__(status_code=502, detail=detail)


class ReasoningTokenLimitExceeded(HTTPException, EmptyModelResponseError):
    """Raised when hidden reasoning consumes the entire output allowance."""

    def __init__(self) -> None:
        super().__init__(status_code=502, detail="输出额度在思考阶段耗尽，请提高最大 Token 后重试")


def is_context_window_error(status_code: int, body: str) -> bool:
    lowered = str(body or "").lower()
    markers = (
        "context length", "context window", "maximum context", "too many tokens",
        "token limit", "prompt is too long", "上下文", "长度超限", "tokens exceed",
    )
    return status_code in {400, 413, 422} and any(marker in lowered for marker in markers)


class ToolsNotSupportedError(HTTPException):
    """Raised when the upstream provider rejects the tools/tool_choice parameters."""

    def __init__(self, detail: str = "\u5f53\u524d\u6a21\u578b\u63a5\u53e3\u4e0d\u652f\u6301\u5de5\u5177\u8c03\u7528") -> None:
        super().__init__(status_code=400, detail=detail)


def is_tools_not_supported_error(status_code: int, body: str) -> bool:
    if status_code not in {400, 404, 422}:
        return False
    lowered = " ".join(str(body or "").lower().split())
    parameter_markers = ("tools", "tool_choice")
    if not any(marker in lowered for marker in parameter_markers):
        return False
    unsupported_markers = (
        "unknown parameter", "unrecognized parameter", "unsupported parameter",
        "unexpected parameter", "extra inputs are not permitted",
        "does not support", "doesn't support", "not supported", "not support",
        "not permitted", "not allowed",
    )
    return any(marker in lowered for marker in unsupported_markers)


def normalize_tool_call(raw: Any, fallback_id: str) -> dict[str, Any] | None:
    """Normalize one OpenAI-style tool_call into {"id", "name", "arguments"} shape."""
    if not isinstance(raw, dict):
        return None
    function = raw.get("function") if isinstance(raw.get("function"), dict) else {}
    name = str(function.get("name") or raw.get("name") or "").strip()
    if not name:
        return None
    arguments: Any = function.get("arguments", raw.get("arguments", {}))
    if isinstance(arguments, str):
        try:
            parsed = json.loads(arguments)
        except (json.JSONDecodeError, ValueError):
            parsed = None
        if isinstance(parsed, dict):
            arguments = parsed
    if not isinstance(arguments, (dict, str)):
        arguments = {}
    return {
        "id": str(raw.get("id") or fallback_id),
        "name": name,
        "arguments": arguments,
    }


def parse_message_tool_calls(message: dict[str, Any]) -> list[dict[str, Any]]:
    raw_calls = message.get("tool_calls") if isinstance(message, dict) else None
    if not isinstance(raw_calls, list):
        return []
    parsed = []
    for position, raw in enumerate(raw_calls):
        normalized = normalize_tool_call(raw, f"call_{position}")
        if normalized is not None:
            parsed.append(normalized)
    return parsed


class ToolCallStreamAggregator:
    """Aggregate OpenAI chat.completions streaming tool_calls deltas by index."""

    def __init__(self) -> None:
        self._slots: dict[int, dict[str, str]] = {}

    def feed(self, raw_calls: Any) -> None:
        if not isinstance(raw_calls, list):
            return
        for position, call in enumerate(raw_calls):
            if not isinstance(call, dict):
                continue
            index = call.get("index")
            slot_key = index if isinstance(index, int) else position
            slot = self._slots.setdefault(slot_key, {"id": "", "name": "", "arguments": ""})
            if call.get("id"):
                slot["id"] = str(call["id"])
            function = call.get("function")
            if isinstance(function, dict):
                if function.get("name"):
                    slot["name"] = str(function["name"])
                if function.get("arguments"):
                    slot["arguments"] += str(function["arguments"])

    def tool_calls(self) -> list[dict[str, Any]]:
        parsed = []
        for slot_key in sorted(self._slots):
            slot = self._slots[slot_key]
            normalized = normalize_tool_call(
                {
                    "id": slot["id"],
                    "function": {"name": slot["name"], "arguments": slot["arguments"]},
                },
                f"call_{slot_key}",
            )
            if normalized is not None:
                parsed.append(normalized)
        return parsed


TEXT_ATTACHMENT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".json", ".jsonl", ".csv", ".tsv", ".xml", ".yaml", ".yml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".scss", ".less", ".java", ".c",
    ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".swift", ".kt", ".kts",
    ".sql", ".sh", ".bat", ".ps1", ".toml", ".ini", ".cfg", ".log",
}
DOCUMENT_ATTACHMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx"}
IMAGE_ATTACHMENT_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
SUPPORTED_ATTACHMENT_EXTENSIONS = (
    TEXT_ATTACHMENT_EXTENSIONS | DOCUMENT_ATTACHMENT_EXTENSIONS | IMAGE_ATTACHMENT_EXTENSIONS
)


class LLMProvider(ABC):
    provider_id: str = ""

    def __init__(
        self,
        config: dict[str, Any],
        *,
        system_prompt_provider: SystemPromptProvider | None = None,
        logger: Any = None,
    ) -> None:
        self.config = config
        self.system_prompt_provider = system_prompt_provider or (lambda: "")
        self.logger = logger

    @classmethod
    @abstractmethod
    def validate_config(cls, config: dict[str, Any]) -> None:
        """Raise ValueError when provider config is incomplete."""

    async def chat(
        self,
        text: str,
        history: list[dict],
        attachments: list[dict[str, Any]] | None = None,
        runtime_context: str = "",
    ) -> str:
        """Compatibility wrapper returning only assistant text."""
        return (await self.chat_result(text, history, attachments, runtime_context)).text

    @abstractmethod
    async def chat_result(
        self,
        text: str,
        history: list[dict],
        attachments: list[dict[str, Any]] | None = None,
        runtime_context: str = "",
    ) -> LLMCallResult:
        """Return assistant text and provider usage metadata."""

    async def stream_result(
        self,
        text: str,
        history: list[dict],
        attachments: list[dict[str, Any]] | None = None,
        runtime_context: str = "",
        *,
        current_user_message_index: int | None = None,
    ) -> AsyncIterator[LLMStreamChunk]:
        """Yield assistant text deltas followed by one terminal result."""
        result = await self.chat_result(text, history, attachments, runtime_context)
        if result.text:
            yield LLMStreamChunk(delta=result.text)
        yield LLMStreamChunk(result=result)

    def system_prompt(self) -> str:
        return str(self.system_prompt_provider() or "")


def require_real_api_key(config: dict[str, Any], provider_label: str) -> str:
    api_key = str(config.get("api_key", "") or "").strip()
    if not has_real_model_api_key(api_key):
        raise ValueError(f"missing {provider_label} api_key")
    return api_key


def require_text_config(config: dict[str, Any], key: str, provider_label: str) -> str:
    value = str(config.get(key, "") or "").strip()
    if not value:
        raise ValueError(f"missing {provider_label} {key}")
    return value


def is_image_attachment(attachment: dict[str, Any]) -> bool:
    return str(attachment.get("mime_type", "")).startswith("image/") or attachment.get("suffix") in IMAGE_ATTACHMENT_EXTENSIONS


def build_vision_content_parts(content: Any, *, supports_vision: bool) -> Any:
    """Build a wire content value for a model message (tool result / user text).

    A ``ModelImagePayload`` stores raw screenshot bytes only in host memory.
    Vision providers encode it into OpenAI-style content parts at the final
    request boundary; non-vision providers receive only its public metadata.

    The raw base64 never sits in the model's text channel: the text part keeps
    only the surrounding prose/note plus a short ``[图片: name]`` marker, and the
    absolute attachment path never reaches the wire at all.
    """
    if isinstance(content, ModelImagePayload):
        if not supports_vision:
            return content.public_text
        parts: list[dict[str, Any]] = [{"type": "text", "text": content.public_text}]
        for image in content.images:
            encoded = base64.b64encode(image.data).decode("ascii")
            parts.append({
                "type": "image_url",
                "image_url": {"url": f"data:{image.mime_type};base64,{encoded}"},
            })
        return parts
    return content


def decode_text_attachment(attachment: dict[str, Any]) -> str:
    data = attachment["data"]
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="缺少 PDF 解析依赖 pypdf，请在 LPP 环境安装") from exc

    reader = PdfReader(BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="缺少 Word 解析依赖 python-docx，请在 LPP 环境安装") from exc

    document = Document(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()


def extract_xlsx_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="缺少 Excel 解析依赖 openpyxl，请在 LPP 环境安装") from exc

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append("\t".join(values))
    workbook.close()
    return "\n".join(lines).strip()


def extract_pptx_text(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="缺少 PPT 解析依赖 python-pptx，请在 LPP 环境安装") from exc

    presentation = Presentation(BytesIO(data))
    lines = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"[Slide {index}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
    return "\n".join(lines).strip()


def extract_attachment_text(attachment: dict[str, Any]) -> str:
    suffix = attachment["suffix"]
    if suffix in TEXT_ATTACHMENT_EXTENSIONS:
        return decode_text_attachment(attachment)
    if suffix == ".pdf":
        return extract_pdf_text(attachment["data"])
    if suffix == ".docx":
        return extract_docx_text(attachment["data"])
    if suffix == ".xlsx":
        return extract_xlsx_text(attachment["data"])
    if suffix == ".pptx":
        return extract_pptx_text(attachment["data"])
    return ""


def build_deepseek_user_text(text: str, attachments: list[dict[str, Any]]) -> str:
    if any(is_image_attachment(attachment) for attachment in attachments):
        raise HTTPException(status_code=400, detail="图片需要切换到支持视觉的模型才能读取")

    sections = [text]
    for attachment in attachments:
        content = extract_attachment_text(attachment)
        if not content:
            content = "[未提取到可读文本]"
        sections.append(
            f"\n\n[附件: {attachment['name']}]\n"
            f"{content[:60000]}"
        )
    return "".join(sections)


def insert_current_user_message(
    history: list[dict],
    current_user_message: dict[str, Any],
    index: int | None = None,
) -> list[dict]:
    """Insert the current turn before tool-loop continuation messages."""
    if index is None:
        return [*history, current_user_message]
    anchor = max(0, min(int(index), len(history)))
    return [*history[:anchor], current_user_message, *history[anchor:]]


def build_responses_input_text(text: str, history: list[dict]) -> str:
    """Legacy text view kept for callers outside the Responses provider."""
    lines = []
    for msg in history:
        role = msg.get("role", "user")
        content = str(msg.get("content", "")).strip()
        if role in ("user", "assistant") and content:
            label = "User" if role == "user" else "Assistant"
            lines.append(f"{label}: {content}")
    lines.append(f"User: {text}")
    return "\n\n".join(lines)


def build_responses_input(
    text: str,
    history: list[dict],
    attachments: list[dict[str, Any]],
    runtime_context: str = "",
    current_user_message_index: int | None = None,
    *,
    supports_vision: bool = True,
) -> list[dict]:
    current_text = text
    if runtime_context.strip():
        current_text = f"{runtime_context.strip()}\n\n[当前用户消息]\n{text}"
    current_content: list[dict[str, Any]] = [{"type": "input_text", "text": current_text}]
    for attachment in attachments:
        if is_image_attachment(attachment):
            current_content.append({
                "type": "input_image",
                "image_url": f"data:{attachment['mime_type']};base64,{attachment['data_base64']}",
            })
        else:
            current_content.append({
                "type": "input_file",
                "filename": attachment["name"],
                "file_data": f"data:{attachment['mime_type']};base64,{attachment['data_base64']}",
            })
    current_user_message = {"role": "user", "content": current_content}

    result: list[dict[str, Any]] = []
    ordered = insert_current_user_message(
        history,
        current_user_message,
        current_user_message_index,
    )
    for msg in ordered:
        role = str(msg.get("role") or "user")
        raw_content = msg.get("content")
        if msg is current_user_message:
            result.append(current_user_message)
            continue
        if role in {"tool", "user"} and isinstance(raw_content, ModelImagePayload):
            content = build_vision_content_parts(raw_content, supports_vision=supports_vision)
            if isinstance(content, list):
                responses_parts: list[dict[str, Any]] = []
                for part in content:
                    if part.get("type") == "image_url":
                        url = part.get("image_url")
                        url = url.get("url") if isinstance(url, dict) else str(url or "")
                        responses_parts.append({"type": "input_image", "image_url": url})
                    else:
                        responses_parts.append({"type": "input_text", "text": str(part.get("text") or "")})
                content = responses_parts
            else:
                content = [{"type": "input_text", "text": str(content or "")}]
            entry: dict[str, Any] = {"role": role, "content": content}
            if role == "tool":
                entry["tool_call_id"] = str(msg.get("tool_call_id") or "")
            result.append(entry)
            continue
        if role == "tool":
            result.append({
                "role": "tool",
                "tool_call_id": str(msg.get("tool_call_id") or ""),
                "content": [{"type": "input_text", "text": str(raw_content or "")}],
            })
            continue
        content = str(raw_content or "").strip()
        if role not in {"user", "assistant", "system"} or not content:
            continue
        result.append({
            "role": "developer" if role == "system" else role,
            "content": [{"type": "input_text", "text": content}],
        })
    return result


def extract_responses_text(data: dict[str, Any]) -> str:
    output_text = data.get("output_text")
    if isinstance(output_text, str) and output_text.strip():
        return output_text.strip()

    chunks: list[str] = []
    for item in data.get("output", []) or []:
        for content in item.get("content", []) or []:
            text = content.get("text")
            if isinstance(text, str):
                chunks.append(text)
    return "\n".join(chunks).strip()


def guess_mime_type(name: str) -> str:
    suffix = Path(name).suffix.lower()
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }.get(suffix, "text/plain")
